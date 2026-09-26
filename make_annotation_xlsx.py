"""Build data/human/annotate.xlsx: a 50-row workbook for writing human examples (movies, news, scores, claims, verify).

Yellow cells are typed or picked from dropdowns; grey cells are lookups. Export with xlsx_to_jsonl.py.

usage: uv run --with openpyxl python make_annotation_xlsx.py [OUT_XLSX]
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
import gen_claims, gen_news

FILMS = {"Solaris Drift": "sci-fi", "Neon Harbour": "cyberpunk", "The Granite Meadow": "drama",
         "Cipher of Lanterns": "mystery", "Velvet Orbit": "romance", "Quiet Ledger": "thriller",
         "Stone River Pact": "western", "Cloud Engine": "adventure", "Midnight Market": "comedy", "Iron Tide": "war"}
LEVELS = ["1 star: hated it, wants a refund", "2 stars: disappointed", "3 stars: decent, mixed feelings",
          "4 stars: liked it", "5 stars: loved it, an instant favourite"]  # as in gen_movie_tasks.py

OUT = sys.argv[1] if len(sys.argv) > 1 else "data/human/annotate.xlsx"
F = "Arial"
HEAD = Font(name=F, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
INPUT = PatternFill("solid", fgColor="FFF2CC")   # yellow: fill in
AUTO = PatternFill("solid", fgColor="EDEDED")    # grey: automatic
EX_FONT = Font(name=F, italic=True, color="7F7F7F")
BODY = Font(name=F)
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

TARGETS = {"Movies": 12, "News": 12, "Scores": 10, "Claims": 8, "Verify": 8}

wb = Workbook()
start = wb.active
start.title = "Start here"
lists = wb.create_sheet("Lists")
bib_ws = wb.create_sheet("Bibliography")

# ---------- Lists (dropdown sources) ----------
films = list(FILMS.items())
cats = list(gen_news.NEWS)
stars = LEVELS
bib = gen_claims.bibliography()
verdicts = gen_claims.VERDICTS
cols = [("Film", [f for f, _ in films]), ("Genre", [g for _, g in films]), ("News category", cats),
        ("Stars", stars), ("Paper key", [b["key"] for b in bib]), ("Verdict", verdicts)]
for c, (name, vals) in enumerate(cols, 1):
    lists.cell(1, c, name).font = HEAD
    lists.cell(1, c).fill = HEAD_FILL
    for r, v in enumerate(vals, 2):
        lists.cell(r, c, v).font = BODY
    lists.column_dimensions[chr(64 + c)].width = 38
RANGE = {"film": f"Lists!$A$2:$A${1 + len(films)}", "genre": f"Lists!$B$2:$B${1 + len(films)}",
         "cat": f"Lists!$C$2:$C${1 + len(cats)}", "stars": f"Lists!$D$2:$D${1 + len(stars)}",
         "key": f"Lists!$E$2:$E${1 + len(bib)}", "verdict": f"Lists!$F$2:$F${1 + len(verdicts)}"}

# ---------- Bibliography (reference for Claims / Verify) ----------
def abstract(p):
    rng = __import__("random").Random(0)
    sents = [gen_claims.cap(gen_claims.phrase(rng, "train", f["dir"], p["intervention"], f["outcome"], f["pct"])) + "."
             for f in p["findings"]]
    return f"Study of {p['intervention']} in {p['population']}. " + " ".join(sents)

bh = ["Paper key", "Full reference", "Topic", "Findings (what the paper says)"]
for c, h in enumerate(bh, 1):
    cell = bib_ws.cell(1, c, h); cell.font = HEAD; cell.fill = HEAD_FILL
for r, p in enumerate(bib, 2):
    for c, v in enumerate([p["key"], p["entry"], p["topic"], abstract(p)], 1):
        cell = bib_ws.cell(r, c, v); cell.font = BODY; cell.alignment = WRAP; cell.border = BOX
for col, w in zip("ABCD", [22, 60, 12, 80]):
    bib_ws.column_dimensions[col].width = w
bib_ws.freeze_panes = "A2"
BIB_KEYS = f"Bibliography!$A$2:$A${1 + len(bib)}"
BIB_REF = f"Bibliography!$B$2:$B${1 + len(bib)}"
BIB_ABS = f"Bibliography!$D$2:$D${1 + len(bib)}"

# ---------- task sheets ----------
# each column: (header, width, kind, validation-range or formula template, note)
SHEETS = {
    "Movies": [
        ("Viewer request (free text: what they want to watch)", 60, "text", None,
         "Describe a film wish in your own words, e.g. a mood, plot or setting. Don't name the title."),
        ("Best film", 24, "list", "film", "Pick the one film that fits the request."),
        ("Genre (auto)", 14, "auto", "=IF(C{r}=\"\",\"\",INDEX(" + RANGE["genre"] + ",MATCH(C{r}," + RANGE["film"] + ",0)))", None),
    ],
    "News": [
        ("Headline (free text)", 70, "text", None, "Write a realistic news headline."),
        ("Category", 18, "list", "cat", "Pick the newspaper section."),
    ],
    "Scores": [
        ("Film watched", 24, "list", "film", "Pick the film the viewer watched."),
        ("What the viewer said afterwards (free text)", 60, "text", None,
         "A short comment, e.g. 'fell asleep halfway' or 'already booked a second viewing'."),
        ("Stars", 36, "list", "stars", "How many stars does that comment imply?"),
    ],
    "Claims": [
        ("Claim (free text, your own wording)", 60, "text", None,
         "Restate ONE finding from the Bibliography sheet in new words. Omitting the population makes it harder (good)."),
        ("Paper it comes from", 24, "list", "key", "Pick the paper key."),
        ("Reference (auto)", 60, "auto", "=IF(C{r}=\"\",\"\",INDEX(" + BIB_REF + ",MATCH(C{r}," + BIB_KEYS + ",0)))", None),
    ],
    "Verify": [
        ("Paper (source)", 24, "list", "key", "Pick a paper; its findings appear on the right."),
        ("Findings of that paper (auto)", 60, "auto", "=IF(B{r}=\"\",\"\",INDEX(" + BIB_ABS + ",MATCH(B{r}," + BIB_KEYS + ",0)))", None),
        ("Claim to check (free text)", 50, "text", None,
         "Write a claim that the paper supports, contradicts (wrong direction or number), or never addresses."),
        ("Verdict", 18, "list", "verdict", "supported / refuted / not enough info"),
    ],
}
EXAMPLES = {
    "Movies": ["a lonely astronaut drifting past Jupiter, something quiet and slow", "Solaris Drift", None],
    "News": ["Regional hospital closes its maternity ward over staff shortages", "health"],
    "Scores": ["Quiet Ledger", "Checked my phone the whole time, couldn't follow the plot", "2 stars: disappointed"],
    "Claims": ["emergency visits went down by about a quarter after a flu booster", "Mancini et al. (2024)", None],
    "Verify": ["Mancini et al. (2024)", None, "the flu booster increased bone density", "refuted"],
}

for name, spec in SHEETS.items():
    ws = wb.create_sheet(name, index=len(wb.sheetnames) - 2)  # task sheets before Lists/Bibliography
    ws.cell(1, 1, "#").font = HEAD
    ws.cell(1, 1).fill = HEAD_FILL
    ws.column_dimensions["A"].width = 10
    n = TARGETS[name]
    last = 2 + n
    for c, (h, w, kind, src, note) in enumerate(spec, 2):
        col = chr(64 + c)
        cell = ws.cell(1, c, h); cell.font = HEAD; cell.fill = HEAD_FILL; cell.alignment = WRAP
        if note:
            cell.comment = Comment(note, "openjev")
        ws.column_dimensions[col].width = w
        if kind == "list":
            dv = DataValidation(type="list", formula1="=" + RANGE[src], allow_blank=True,
                                showErrorMessage=True, errorTitle="Pick from the list",
                                error="Choose a value from the dropdown.")
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{last}")
    ws.row_dimensions[1].height = 32
    # row 2: example
    ws.cell(2, 1, "EXAMPLE").font = EX_FONT
    for c, (h, w, kind, src, note) in enumerate(spec, 2):
        cell = ws.cell(2, c)
        cell.value = src.format(r=2) if kind == "auto" else EXAMPLES[name][c - 2]
        cell.font = EX_FONT; cell.alignment = WRAP; cell.border = BOX
    # rows 3..: input rows
    for i in range(1, n + 1):
        r = 2 + i
        ws.cell(r, 1, i).font = BODY
        for c, (h, w, kind, src, note) in enumerate(spec, 2):
            cell = ws.cell(r, c)
            cell.font = BODY; cell.alignment = WRAP; cell.border = BOX
            if kind == "auto":
                cell.value = src.format(r=r); cell.fill = AUTO
            else:
                cell.fill = INPUT
        ws.row_dimensions[r].height = 30
    ws.freeze_panes = "B3"

# ---------- Start here ----------
start.column_dimensions["A"].width = 16
start.column_dimensions["B"].width = 14
start.column_dimensions["C"].width = 14
start.column_dimensions["D"].width = 70
start["A1"] = "openjev — human examples (50 rows)"
start["A1"].font = Font(name=F, bold=True, size=14)
lines = [
    "Fill the YELLOW cells on each task sheet. Grey cells fill themselves; don't edit them.",
    "Free-text columns: write in your own words, one short sentence. Answer columns: use the dropdown.",
    "Row 2 on each sheet is an EXAMPLE (grey italic) and is not exported.",
    "Claims and Verify use the invented papers on the Bibliography sheet; read a paper's findings first.",
    "Hover a column header for a hint. Partly filled rows are skipped on export.",
]
for i, t in enumerate(lines, 3):
    start.cell(i, 1, t).font = BODY
hr = 3 + len(lines) + 1
for c, h in enumerate(["Sheet", "Filled", "Target", "What you do"], 1):
    cell = start.cell(hr, c, h); cell.font = HEAD; cell.fill = HEAD_FILL
what = {"Movies": "Write a film wish, pick the film that fits",
        "News": "Write a headline, pick its section",
        "Scores": "Pick a film, write the viewer's reaction, pick the stars",
        "Claims": "Rephrase a paper's finding, pick the paper",
        "Verify": "Pick a paper, write a claim, say if the paper supports it"}
answer_col = {"Movies": "C", "News": "C", "Scores": "D", "Claims": "C", "Verify": "E"}
text_col = {"Movies": "B", "News": "B", "Scores": "C", "Claims": "B", "Verify": "D"}
for i, name in enumerate(SHEETS, hr + 1):
    last = 2 + TARGETS[name]
    t, a = text_col[name], answer_col[name]
    start.cell(i, 1, name).font = BODY
    need = [t, a] + (["B"] if name in ("Scores", "Verify") else [])  # every column the export requires
    start.cell(i, 2, "=COUNTIFS(" + ",".join(f"{name}!{c}3:{c}{last},\"<>\"" for c in need) + ")").font = BODY
    start.cell(i, 3, TARGETS[name]).font = BODY
    start.cell(i, 4, what[name]).font = BODY
tot = hr + 1 + len(SHEETS)
start.cell(tot, 1, "Total").font = Font(name=F, bold=True)
start.cell(tot, 2, f"=SUM(B{hr + 1}:B{tot - 1})").font = Font(name=F, bold=True)
start.cell(tot, 3, f"=SUM(C{hr + 1}:C{tot - 1})").font = Font(name=F, bold=True)
start.cell(tot + 2, 1, "Export: .venv/bin/python xlsx_to_jsonl.py data/human/annotate.xlsx  →  data/human/<task>.jsonl").font = EX_FONT
lg = tot + 4
start.cell(lg, 1, "Legend").font = Font(name=F, bold=True)
start.cell(lg + 1, 1, "fill in").fill = INPUT
start.cell(lg + 2, 1, "automatic").fill = AUTO
for r in (lg + 1, lg + 2):
    start.cell(r, 1).font = BODY

import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print(OUT)
