"""Export the human-written rows of data/human/annotate.xlsx to openjev JSONL.

Reads only the typed and dropdown cells (no formulas), skips the EXAMPLE row and any row with an
empty required cell, and writes one file per task, in the same format as the synthetic sets:
movies (request -> 1 of 10 films), news (headline -> 8 sections), scores (film + comment -> 1-5
stars), claims_keys / claims_full (claim -> paper, as keys or full references, with the same-topic
papers among the 8 options), verify (paper findings + claim -> supported/refuted/not enough info).

usage: uv run --with openpyxl python xlsx_to_jsonl.py [XLSX] [OUT_DIR] [SEED]
"""
import json, os, random, sys
from openpyxl import load_workbook
import gen_claims, gen_news

FILMS = ["Solaris Drift", "Neon Harbour", "The Granite Meadow", "Cipher of Lanterns", "Velvet Orbit",
         "Quiet Ledger", "Stone River Pact", "Cloud Engine", "Midnight Market", "Iron Tide"]
LEVELS = ["1 star: hated it, wants a refund", "2 stars: disappointed", "3 stars: decent, mixed feelings",
          "4 stars: liked it", "5 stars: loved it, an instant favourite"]  # as in gen_movie_tasks.py


def rows(wb, sheet, cols):
    """Yield {name: value} for filled rows; cols maps name -> column letter."""
    ws = wb[sheet]
    for r in range(3, ws.max_row + 1):
        vals = {k: ws[f"{c}{r}"].value for k, c in cols.items()}
        vals = {k: v.strip() if isinstance(v, str) else v for k, v in vals.items()}
        if all(vals.values()):
            yield vals


def shuffled(rng, opts, target):
    opts = opts[:]
    rng.shuffle(opts)
    return opts, opts.index(target)


def export(wb, rng):
    bib = gen_claims.bibliography()
    by_key = {p["key"]: p for p in bib}
    findings = {r[0].value: r[3].value for r in wb["Bibliography"].iter_rows(min_row=2) if r[0].value}
    out = {k: [] for k in ["movies", "news", "scores", "claims_keys", "claims_full", "verify"]}

    for v in rows(wb, "Movies", {"req": "B", "film": "C"}):
        opts, y = shuffled(rng, FILMS, v["film"])
        out["movies"].append({"context": f"A viewer asks for {v['req'].rstrip('.')}. Recommend exactly one film from the list.",
                              "options": opts, "label": y, "film": v["film"]})
    for v in rows(wb, "News", {"h": "B", "cat": "C"}):
        opts, y = shuffled(rng, list(gen_news.NEWS), v["cat"])
        out["news"].append({"context": f"Classify this headline: {v['h'].rstrip('.')}.", "options": opts, "label": y,
                            "category": v["cat"]})
    for v in rows(wb, "Scores", {"film": "B", "said": "C", "stars": "D"}):
        out["scores"].append({"context": f"Film watched: {v['film']}. The viewer said afterwards: {v['said']} Rate the viewing.",
                              "options": LEVELS, "label": LEVELS.index(v["stars"])})
    for v in rows(wb, "Claims", {"claim": "B", "key": "C"}):
        t = by_key[v["key"]]
        sib = [p for p in bib if p["topic"] == t["topic"] and p is not t]
        papers = [t] + sib + rng.sample([p for p in bib if p["topic"] != t["topic"]], 5)
        rng.shuffle(papers)
        ctx = f"Claim: {v['claim'].rstrip('.')}. Which reference supports this claim?"
        for name, field in (("claims_keys", "key"), ("claims_full", "entry")):
            out[name].append({"context": ctx, "options": [p[field] for p in papers], "label": papers.index(t),
                              "key": t["key"], "claim": v["claim"]})
    for v in rows(wb, "Verify", {"key": "B", "claim": "D", "verdict": "E"}):
        opts, y = shuffled(rng, gen_claims.VERDICTS, v["verdict"])
        out["verify"].append({"context": f"Source abstract: {findings[v['key']]}\nClaim: {v['claim'].rstrip('.')}.\n"
                                         "Does the source support the claim?",
                              "options": opts, "label": y, "verdict": v["verdict"], "key": v["key"]})
    return out


if __name__ == "__main__":
    xlsx = sys.argv[1] if len(sys.argv) > 1 else "data/human/annotate.xlsx"
    out_dir = sys.argv[2] if len(sys.argv) > 2 else "data/human"
    rng = random.Random(int(sys.argv[3]) if len(sys.argv) > 3 else 5)
    data = export(load_workbook(xlsx), rng)
    os.makedirs(out_dir, exist_ok=True)
    for name, rs in data.items():
        with open(f"{out_dir}/{name}.jsonl", "w") as f:
            for r in rs:
                f.write(json.dumps(r, ensure_ascii=False) + "\n")
        print(f"{name}: {len(rs)} rows")
